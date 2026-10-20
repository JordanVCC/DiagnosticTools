"""Mirror knowledge source folders into agent/knowledge/sources as Markdown.

Run with the agent environment:
    .venv\\Scripts\\python.exe sync_knowledge.py
Unchanged files are skipped; outputs whose source was removed are deleted.
"""

import json
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

AGENT_DIR = Path(__file__).resolve().parent
SOURCES_FILE = AGENT_DIR / "knowledge_sources.txt"
OUTPUT_DIR = AGENT_DIR / "knowledge" / "sources"
MANIFEST = OUTPUT_DIR / ".sync_manifest.json"
TEXT_SUFFIXES = {".md", ".txt", ".csv", ".json", ".yaml", ".yml"}
# Pages with more vector paths than this (e.g. component placement drawings) stall table detection.
MAX_PAGE_DRAWINGS = 20000
# Pages whose Markdown keeps less of the PDF's words than this get the page's raw words appended
# (schematic pin/ball labels such as AH41 are otherwise dropped or glued together).
MIN_PAGE_COVERAGE = 0.85
COVERAGE_WORD = re.compile(r"[A-Za-z0-9_]{4,}")
# PDFs with less extractable text than this per page are scans or outlined text and need OCR.
MIN_TEXT_CHARS_PER_PAGE = 20
OCR_PYTHON = Path(r"C:\pmt\Scripts\python.exe")
OCR_TOOL = AGENT_DIR.parents[1] / "pdf_to_markdown_tool" / "pdf_to_markdown.py"
OCR_TIMEOUT_S = 4 * 60 * 60
# Bump when converter output changes so existing outputs are regenerated.
CONVERTER_VERSION = 2


def read_sources() -> list[Path]:
    lines = SOURCES_FILE.read_text(encoding="utf-8-sig").splitlines()
    return [Path(line.strip().strip('"')) for line in lines if line.strip() and not line.lstrip().startswith("#")]


def output_path(source_root: Path, path: Path) -> Path:
    rel = path.relative_to(source_root)
    suffix = path.suffix.lower()
    name = rel.name if suffix in TEXT_SUFFIXES else rel.stem + ".md"
    return OUTPUT_DIR / source_root.name / rel.parent / name


def header(path: Path) -> str:
    return f"# {path.stem}\n\nSource: {path}\n\n"


def coverage_words(text: str) -> set[str]:
    return {w.lower() for w in COVERAGE_WORD.findall(text)}


def page_word_text(page) -> str:
    """Page text rebuilt from PDF words, one output line per text line, so labels stay separate."""
    lines: dict[tuple[int, int], list[str]] = {}
    for *_, word, block, line, _ in page.get_text("words"):
        lines.setdefault((block, line), []).append(word)
    return "\n".join(" ".join(words) for words in lines.values())


def convert_pdf_ocr(path: Path, out: Path) -> None:
    """OCR a PDF with pdf_to_markdown_tool (marker-pdf) in its own environment."""
    if not OCR_PYTHON.exists() or not OCR_TOOL.exists():
        raise RuntimeError(f"no extractable text; OCR needs {OCR_PYTHON} and {OCR_TOOL}")
    with tempfile.TemporaryDirectory() as tmp:
        tmp_out = Path(tmp) / "ocr.md"
        subprocess.run([str(OCR_PYTHON), str(OCR_TOOL), str(path), str(tmp_out)],
                       check=True, timeout=OCR_TIMEOUT_S)
        out.write_text(header(path) + tmp_out.read_text(encoding="utf-8"), encoding="utf-8")
    print("  converted with OCR (marker-pdf)", flush=True)


def convert_pdf(path: Path, out: Path) -> None:
    import pymupdf
    import pymupdf4llm

    doc = pymupdf.open(str(path))
    page_texts = [page.get_text() for page in doc]
    if sum(len(t.strip()) for t in page_texts) < MIN_TEXT_CHARS_PER_PAGE * len(doc):
        doc.close()
        convert_pdf_ocr(path, out)
        return

    dense = {i for i, page in enumerate(doc) if len(page.get_drawings()) > MAX_PAGE_DRAWINGS}
    normal = [i for i in range(len(doc)) if i not in dense]
    markdown = {}
    if normal:
        for chunk in pymupdf4llm.to_markdown(doc, pages=normal, page_chunks=True, show_progress=True):
            markdown[chunk["metadata"]["page_number"] - 1] = chunk["text"]

    parts, source_words, kept_words, weak = [header(path)], set(), set(), 0
    for i, page in enumerate(doc):
        words = coverage_words(page_texts[i])
        if i in dense:
            text = f"\n## Page {i + 1} (drawing, text only)\n\n{page_word_text(page)}\n"
        else:
            text = markdown.get(i, "")
            if words and len(words & coverage_words(text)) / len(words) < MIN_PAGE_COVERAGE:
                weak += 1
                text += f"\n\n### Page {i + 1} labels and values (raw PDF text)\n\n{page_word_text(page)}\n"
        parts.append(text)
        source_words |= words
        kept_words |= words & coverage_words(text)
    doc.close()
    out.write_text("".join(parts), encoding="utf-8")
    coverage = len(kept_words) / len(source_words) if source_words else 1.0
    print(f"  text coverage {coverage:.1%}; {weak} weak page(s) supplemented, {len(dense)} drawing page(s)",
          flush=True)


def convert_xlsx(path: Path, out: Path) -> None:
    from openpyxl import load_workbook

    def cell(value) -> str:
        return "" if value is None else str(value).replace("|", "\\|").replace("\n", " ").strip()

    parts = [header(path)]
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = [[cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        used = [c for c in range(width) if any(row[c] for row in rows)]
        rows = [[row[c] for c in used] for row in rows]
        width = len(used)
        parts.append(f"## {sheet.title}\n\n| " + " | ".join(rows[0]) + " |\n|" + "---|" * width + "\n")
        parts.extend("| " + " | ".join(row) + " |\n" for row in rows[1:])
        parts.append("\n")
    workbook.close()
    out.write_text("".join(parts), encoding="utf-8")


def convert_url(path: Path, out: Path) -> None:
    url = next((line[4:].strip() for line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
                if line.upper().startswith("URL=")), "")
    out.write_text(header(path) + f"Link: {url}\n", encoding="utf-8")


def copy_text(path: Path, out: Path) -> None:
    out.write_bytes(path.read_bytes())


CONVERTERS = {".pdf": convert_pdf, ".xlsx": convert_xlsx, ".xlsm": convert_xlsx, ".url": convert_url}
CONVERTERS.update({suffix: copy_text for suffix in TEXT_SUFFIXES})


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8")) if MANIFEST.exists() else {}
    seen, converted, skipped, failed, unsupported = set(), 0, 0, [], []

    for root in read_sources():
        if not root.is_dir():
            failed.append((str(root), "source folder not found"))
            continue
        for path in sorted(p for p in root.rglob("*") if p.is_file()):
            converter = CONVERTERS.get(path.suffix.lower())
            if not converter:
                unsupported.append(str(path))
                continue
            key = str(path)
            seen.add(key)
            out = output_path(root, path)
            stat = path.stat()
            signature = {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "output": str(out),
                         "converter": CONVERTER_VERSION}
            if manifest.get(key) == signature and out.exists():
                skipped += 1
                continue

            print(f"Converting {path.relative_to(root.parent)} ...", flush=True)
            started = time.time()
            out.parent.mkdir(parents=True, exist_ok=True)
            try:
                converter(path, out)
            except Exception as exc:
                failed.append((key, f"{type(exc).__name__}: {exc}"))
                print(f"  FAILED: {exc}", flush=True)
                continue
            manifest[key] = signature
            MANIFEST.write_text(json.dumps(manifest, indent=1), encoding="utf-8")
            converted += 1
            print(f"  done in {time.time() - started:.0f} s -> {out.relative_to(OUTPUT_DIR)}", flush=True)

    removed = 0
    for key in [k for k in manifest if k not in seen]:
        Path(manifest.pop(key)["output"]).unlink(missing_ok=True)
        removed += 1
    MANIFEST.write_text(json.dumps(manifest, indent=1), encoding="utf-8")

    print(f"\nConverted {converted}, unchanged {skipped}, removed {removed}, failed {len(failed)}, "
          f"unsupported {len(unsupported)}.")
    for name, reason in failed:
        print(f"  FAILED {name}: {reason}")
    for name in unsupported:
        print(f"  UNSUPPORTED {name}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
